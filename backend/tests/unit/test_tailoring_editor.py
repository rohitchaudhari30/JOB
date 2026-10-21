import hashlib

from docx import Document

from app.domain.resume_tailoring.plan import ProposedEdit
from app.infrastructure.documents.docx_editor import apply_edits
from app.infrastructure.documents.docx_map import build_docx_map
from app.infrastructure.documents.fidelity import FAIL, PASS, UNAVAILABLE, validate_tailored_docx
from app.infrastructure.documents.renderer import renderer_available
from tests.support.resume_docs import make_resume


def _block(doc, start):
    return next(b for b in doc.blocks if b.text.startswith(start))


def _runs(path, start):
    p = next(p for p in Document(str(path)).paragraphs if p.text.startswith(start))
    return [(r.text, r.bold, r.font.size, r.font.name) for r in p.runs]


def test_rewording_keeps_run_formatting_and_leaves_the_master_untouched(tmp_path):
    master = make_resume(tmp_path / "m.docx")
    before_hash = hashlib.sha256(master.read_bytes()).hexdigest()
    doc = build_docx_map(str(master))
    skills = _block(doc, "Languages")
    edit = ProposedEdit("e1", "REWORD", skills.id, skills.text, "Languages: SQL, Python, Bash scripting")
    out = tmp_path / "t.docx"
    result = apply_edits(str(master), str(out), [edit])
    assert result.applied == ["e1"] and not result.skipped
    # Bold label run is exactly as it was; only the second run's characters changed.
    assert _runs(out, "Languages")[0] == _runs(master, "Languages")[0]
    assert _runs(out, "Languages")[1][0] == "SQL, Python, Bash scripting"
    assert hashlib.sha256(master.read_bytes()).hexdigest() == before_hash


def test_edit_spanning_two_formatted_runs_keeps_each_runs_formatting(tmp_path):
    master = make_resume(tmp_path / "m.docx")
    doc = build_docx_map(str(master))
    skills = _block(doc, "Languages")
    edit = ProposedEdit("e1", "REWORD", skills.id, skills.text, "Core languages: Python, SQL, Bash scripting")
    out = tmp_path / "t.docx"
    assert apply_edits(str(master), str(out), [edit]).applied == ["e1"]
    runs = _runs(out, "Core languages")
    assert runs[0][1] is True and runs[1][1] is None  # label still bold, list still plain


def test_reorder_swaps_two_bullets_without_changing_their_formatting(tmp_path):
    master = make_resume(tmp_path / "m.docx")
    doc = build_docx_map(str(master))
    a, b = _block(doc, "Built Python"), _block(doc, "Reduced report")
    out = tmp_path / "t.docx"
    result = apply_edits(str(master), str(out), [ProposedEdit("r1", "REORDER", a.id, target_block_id=b.id)])
    assert result.applied == ["r1"]
    texts = [p.text for p in Document(str(out)).paragraphs if p.style.name == "List Bullet"]
    assert texts[0].startswith("Reduced report") and texts[1].startswith("Built Python")
    report = validate_tailored_docx(str(master), str(out), [ProposedEdit("r1", "REORDER", a.id, target_block_id=b.id)])
    assert report.passed, report.to_dict()


def test_two_edits_on_the_same_line_apply_only_once(tmp_path):
    master = make_resume(tmp_path / "m.docx")
    doc = build_docx_map(str(master))
    b = _block(doc, "Built Python")
    e1 = ProposedEdit("e1", "REWORD", b.id, b.text, "Built scalable Python data pipelines that load daily reports into PostgreSQL.")
    e2 = ProposedEdit("e2", "REWORD", b.id, b.text, "Built robust Python data pipelines that load daily reports into PostgreSQL.")
    result = apply_edits(str(master), str(tmp_path / "t.docx"), [e1, e2])
    assert result.applied == ["e1"] and result.skipped[0][0] == "e2"


def test_validation_passes_for_a_clean_edit_and_reports_every_check(tmp_path):
    master = make_resume(tmp_path / "m.docx")
    doc = build_docx_map(str(master))
    b = _block(doc, "Built Python")
    edit = ProposedEdit("e1", "REWORD", b.id, b.text, "Built scalable Python data pipelines that load daily reports into PostgreSQL.")
    out = tmp_path / "t.docx"
    apply_edits(str(master), str(out), [edit])
    report = validate_tailored_docx(str(master), str(out), [edit])
    names = {c.name: c.status for c in report.checks}
    assert report.passed and names["structure"] == PASS and names["wording"] == PASS and names["protected_facts"] == PASS
    assert names["layout"] == (PASS if renderer_available() else UNAVAILABLE)


def test_validation_catches_a_changed_line_that_was_not_approved(tmp_path):
    master = make_resume(tmp_path / "m.docx")
    out = tmp_path / "t.docx"
    doc = Document(str(master))
    next(p for p in doc.paragraphs if p.text.startswith("Led a team")).runs[0].text = "Led a team of 40 engineers and reviewed all production releases."
    doc.save(str(out))
    report = validate_tailored_docx(str(master), str(out), [])
    statuses = {c.name: c.status for c in report.checks}
    assert not report.passed and statuses["structure"] == FAIL and statuses["wording"] == FAIL


def test_validation_catches_a_changed_contact_detail(tmp_path):
    master = make_resume(tmp_path / "m.docx")
    out = tmp_path / "t.docx"
    doc = Document(str(master))
    p = next(p for p in doc.paragraphs if "asha.verma@" in p.text)
    p.runs[0].text = p.text.replace("asha.verma@example.com", "someone.else@example.com")
    doc.save(str(out))
    report = validate_tailored_docx(str(master), str(out), [])
    assert {c.name: c.status for c in report.checks}["protected_facts"] == FAIL


def test_validation_catches_overflow_to_an_extra_page_when_a_renderer_exists(tmp_path):
    if not renderer_available():
        return
    master = make_resume(tmp_path / "m.docx", extra_bullets=22)
    doc = build_docx_map(str(master))
    b = _block(doc, "Built Python")
    long_text = b.text + " " + ("Also supported many other reporting needs for the whole company. " * 25)
    edit = ProposedEdit("e1", "REWORD", b.id, b.text, long_text)
    out = tmp_path / "t.docx"
    apply_edits(str(master), str(out), [edit])  # editor alone does not judge length; the checks do
    report = validate_tailored_docx(str(master), str(out), [edit])
    layout = next(c for c in report.checks if c.name == "layout")
    assert layout.status == FAIL
