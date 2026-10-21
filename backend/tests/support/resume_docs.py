"""Builds real DOCX resumes for tests (python-docx writes genuine Word files; nothing is faked)."""
from docx import Document
from docx.shared import Pt


def make_resume(path, *, table_skills=False, header_text=None, bold_skills=True, extra_bullets=0):
    doc = Document()
    if header_text:
        doc.sections[0].header.paragraphs[0].text = header_text
    name = doc.add_paragraph()
    run = name.add_run("Asha Verma")
    run.bold = True
    run.font.size = Pt(18)
    doc.add_paragraph("Backend Engineer")
    doc.add_paragraph("asha.verma@example.com | +91 98765 43210 | linkedin.com/in/ashaverma")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph("Backend engineer with 6 years of experience building Python services and data pipelines.")
    doc.add_heading("Skills", level=2)
    if table_skills:
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Languages"
        table.cell(0, 1).text = "Python, SQL, Bash scripting"
    else:
        p = doc.add_paragraph()
        r1 = p.add_run("Languages: ")
        r1.bold = bold_skills
        p.add_run("Python, SQL, Bash scripting")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Senior Engineer, Northwind Labs (2021 - 2025)")
    doc.add_paragraph("Built Python data pipelines that load daily reports into PostgreSQL.", style="List Bullet")
    doc.add_paragraph("Reduced report generation time by 40% by caching results in Redis.", style="List Bullet")
    doc.add_paragraph("Led a team of 4 engineers and reviewed all production releases.", style="List Bullet")
    for i in range(extra_bullets):
        doc.add_paragraph(f"Maintained internal tooling number {i} used by the support team.", style="List Bullet")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.E. Computer Engineering, Pune University (2014 - 2018)")
    doc.save(path)
    return path
